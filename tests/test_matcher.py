import json
from io import BytesIO

from docx import Document
from src.jd_resume_agent.agent_pipeline import get_agent_steps, run_agent_pipeline
from src.jd_resume_agent.llm_advisor import (
    analyze_with_optional_llm,
    build_llm_prompt,
    merge_llm_advice,
    parse_llm_json,
)
from src.jd_resume_agent.matcher import analyze_resume_match, extract_skills
from src.jd_resume_agent.parser import read_docx_file, read_pdf_file, read_txt_file


def test_extract_skills_finds_known_keywords():
    text = "熟悉 Python、Streamlit、RAG 和 Git。"

    skills = extract_skills(text)

    assert "python" in skills
    assert "streamlit" in skills
    assert "rag" in skills
    assert "git" in skills


def test_analyze_resume_match_scores_by_jd_skills():
    jd = "要求熟悉 Python、Streamlit、FastAPI、pytest。"
    resume = "项目使用 Python 和 Streamlit 搭建问答 Demo。"

    result = analyze_resume_match(jd, resume)

    assert result["status"] == "ok"
    assert result["match_score"] == 50
    assert result["matched_skills"] == ["python", "streamlit"]
    assert result["missing_skills"] == ["fastapi", "pytest"]


def test_analyze_resume_match_has_coach_fields():
    result = analyze_resume_match("Python FastAPI pytest", "Python")

    assert result["analysis_mode"] == "rule_based"
    assert result["resume_strengths"]
    assert result["gap_analysis"][0]["skill"] == "fastapi"
    assert result["rewrite_suggestions"]
    assert result["next_steps"]


def test_agent_pipeline_adds_step_trace():
    result = run_agent_pipeline("Python FastAPI", "Python")

    assert result["analysis_mode"] == "rule_based"
    assert result["agent_steps"]
    assert result["agent_steps"][0]["name"] == "input_validation"


def test_get_agent_steps_is_structured():
    steps = get_agent_steps()

    assert len(steps) >= 4
    assert {"name", "description"} == set(steps[0].keys())


def test_analyze_resume_match_handles_missing_input():
    result = analyze_resume_match("", "Python")

    assert result["status"] == "missing_input"
    assert result["match_score"] == 0
    assert result["next_steps"] == ["输入两段文本后，再点击分析按钮。"]


def test_analyze_resume_match_result_is_json_serializable():
    result = analyze_resume_match("Python FastAPI", "Python")

    json_text = json.dumps(result, ensure_ascii=False)

    assert "match_score" in json_text
    assert "gap_analysis" in json_text


def test_build_llm_prompt_contains_required_json_fields():
    rule_result = analyze_resume_match("Python", "Python")

    prompt = build_llm_prompt("Python", "Python", rule_result)

    assert "resume_strengths" in prompt
    assert "gap_analysis" in prompt
    assert "只输出 JSON" in prompt


def test_parse_llm_json_validates_required_fields():
    content = json.dumps(
        {
            "resume_strengths": ["优势"],
            "gap_analysis": [],
            "rewrite_suggestions": ["建议"],
            "next_steps": ["下一步"],
        },
        ensure_ascii=False,
    )

    data = parse_llm_json(content)

    assert data["resume_strengths"] == ["优势"]


def test_merge_llm_advice_marks_llm_enhanced():
    rule_result = analyze_resume_match("Python FastAPI", "Python")
    llm_advice = {
        "resume_strengths": ["LLM 优势"],
        "gap_analysis": [],
        "rewrite_suggestions": ["LLM 建议"],
        "next_steps": ["LLM 下一步"],
    }

    result = merge_llm_advice(rule_result, llm_advice)

    assert result["analysis_mode"] == "llm_enhanced"
    assert result["llm_status"] == "ok"
    assert result["resume_strengths"] == ["LLM 优势"]


def test_analyze_with_optional_llm_skips_without_api_key(monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    rule_result = analyze_resume_match("Python", "Python")

    result = analyze_with_optional_llm("Python", "Python", rule_result)

    assert result["analysis_mode"] == "rule_based"
    assert result["llm_status"] == "skipped"


def test_read_txt_file_supports_utf8_bytes():
    text = read_txt_file("熟悉 Python 和 Streamlit".encode("utf-8"))

    assert "Python" in text


def test_read_txt_file_supports_gbk_bytes():
    text = read_txt_file("熟悉 Python 和 Streamlit".encode("gbk"))

    assert "Python" in text


def test_read_docx_file_extracts_paragraph_text():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("熟悉 Python 和 FastAPI")
    document.add_paragraph("做过 AI Agent 项目")
    document.save(buffer)

    text = read_docx_file(buffer.getvalue())

    assert "Python" in text
    assert "AI Agent" in text


def test_read_pdf_file_reports_missing_dependency_when_needed(monkeypatch):
    import builtins

    original_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "pypdf":
            raise ImportError("missing pypdf")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    try:
        read_pdf_file(b"not a real pdf")
    except RuntimeError as exc:
        assert "pypdf" in str(exc)
    else:
        raise AssertionError("Expected RuntimeError when pypdf is missing")


def test_read_pdf_file_reports_empty_extraction(monkeypatch):
    import sys
    import types

    class FakePage:
        def extract_text(self):
            return ""

    class FakeReader:
        def __init__(self, _file_obj):
            self.pages = [FakePage()]

    fake_pypdf = types.SimpleNamespace(PdfReader=FakeReader)
    monkeypatch.setitem(sys.modules, "pypdf", fake_pypdf)

    try:
        read_pdf_file(b"%PDF fake")
    except RuntimeError as exc:
        assert "未提取到文字" in str(exc)
    else:
        raise AssertionError("Expected RuntimeError when PDF text is empty")
